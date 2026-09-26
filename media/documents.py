"""Document text extraction: PDF (with scanned-page fallback), DOCX, PPTX, XLSX/CSV, text/code."""

from __future__ import annotations

import csv
import io
from dataclasses import asdict, dataclass, field
from pathlib import Path

from media.detect import CODE_EXT

MAX_TABLE_ROWS = 500
SCANNED_PAGE_MIN_CHARS = 25


@dataclass
class DocSection:
    label: str          # "page 3", "slide 2", "sheet Revenue", ...
    text: str
    image_path: str | None = None  # rendered page image for scanned pages


@dataclass
class DocResult:
    format: str
    sections: list[DocSection] = field(default_factory=list)
    meta: dict = field(default_factory=dict)

    @property
    def text(self) -> str:
        return "\n\n".join(f"--- {s.label} ---\n{s.text}" if s.label else s.text
                           for s in self.sections if s.text.strip())

    def to_dict(self) -> dict:
        return {"format": self.format, "meta": self.meta,
                "sections": [asdict(s) for s in self.sections]}

    @classmethod
    def from_dict(cls, d: dict) -> DocResult:
        return cls(format=d["format"], meta=d.get("meta", {}),
                   sections=[DocSection(**s) for s in d.get("sections", [])])


def extract_document(path: Path, filename: str, render_dir: Path | None = None,
                     max_rendered_pages: int = 20, render_dpi: int = 150) -> DocResult:
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        return _pdf(path, render_dir, max_rendered_pages, render_dpi)
    if ext == ".docx":
        return _docx(path)
    if ext == ".pptx":
        return _pptx(path)
    if ext in (".xlsx", ".xlsm"):
        return _xlsx(path)
    if ext in (".csv", ".tsv"):
        return _csv(path, "\t" if ext == ".tsv" else ",")
    return _text(path, ext)


def _pdf(path: Path, render_dir: Path | None, max_rendered: int, dpi: int) -> DocResult:
    import fitz  # pymupdf

    res = DocResult(format="pdf")
    rendered = 0
    with fitz.open(path) as doc:
        res.meta = {"pages": doc.page_count, "title": (doc.metadata or {}).get("title") or None}
        for i, page in enumerate(doc, start=1):
            text = page.get_text("text").strip()
            section = DocSection(label=f"page {i}", text=text)
            if len(text) < SCANNED_PAGE_MIN_CHARS and render_dir is not None and rendered < max_rendered:
                render_dir.mkdir(parents=True, exist_ok=True)
                out = render_dir / f"page_{i:04d}.png"
                page.get_pixmap(dpi=dpi).save(out)
                section.image_path = str(out)
                rendered += 1
            res.sections.append(section)
    res.meta["scanned_pages"] = rendered
    return res


def _docx(path: Path) -> DocResult:
    import docx

    d = docx.Document(str(path))
    parts: list[str] = []
    for block in d.element.body.iterchildren():
        tag = block.tag.rsplit("}", 1)[-1]
        if tag == "p":
            para = docx.text.paragraph.Paragraph(block, d)
            style = (para.style.name or "").lower() if para.style is not None else ""
            text = para.text.strip()
            if not text:
                continue
            if style.startswith("heading"):
                level = "".join(ch for ch in style if ch.isdigit()) or "1"
                text = "#" * min(int(level), 6) + " " + text
            elif "list" in style:
                text = "- " + text
            parts.append(text)
        elif tag == "tbl":
            table = docx.table.Table(block, d)
            rows = [[c.text.strip() for c in r.cells] for r in table.rows]
            parts.append(_md_table(rows))
    return DocResult(format="docx", sections=[DocSection(label="", text="\n\n".join(parts))])


def _pptx(path: Path) -> DocResult:
    from pptx import Presentation

    prs = Presentation(str(path))
    res = DocResult(format="pptx", meta={"slides": len(prs.slides)})
    for i, slide in enumerate(prs.slides, start=1):
        lines = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                lines += [p.text for p in shape.text_frame.paragraphs if p.text.strip()]
            if getattr(shape, "has_table", False) and shape.has_table:
                rows = [[c.text for c in r.cells] for r in shape.table.rows]
                lines.append(_md_table(rows))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                lines.append(f"[speaker notes] {notes}")
        res.sections.append(DocSection(label=f"slide {i}", text="\n".join(lines)))
    return res


def _xlsx(path: Path) -> DocResult:
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    res = DocResult(format="xlsx", meta={"sheets": wb.sheetnames})
    for ws in wb.worksheets:
        rows = []
        for r in ws.iter_rows(values_only=True):
            if any(v is not None for v in r):
                rows.append(["" if v is None else str(v) for v in r])
            if len(rows) > MAX_TABLE_ROWS:
                break
        text = _md_table(rows[:MAX_TABLE_ROWS])
        if len(rows) > MAX_TABLE_ROWS:
            text += f"\n… truncated after {MAX_TABLE_ROWS} rows"
        res.sections.append(DocSection(label=f"sheet {ws.title}", text=text))
    wb.close()
    return res


def _csv(path: Path, delimiter: str) -> DocResult:
    raw = path.read_bytes().decode("utf-8", errors="replace")
    try:
        delimiter = csv.Sniffer().sniff(raw[:4096], delimiters=",;\t|").delimiter
    except csv.Error:
        pass
    rows = []
    for i, r in enumerate(csv.reader(io.StringIO(raw), delimiter=delimiter)):
        if i > MAX_TABLE_ROWS:
            break
        rows.append(r)
    text = _md_table(rows[:MAX_TABLE_ROWS])
    total = raw.count("\n")
    if total > MAX_TABLE_ROWS:
        text += f"\n… {total} rows total; showing the first {MAX_TABLE_ROWS}"
    return DocResult(format="csv", sections=[DocSection(label="", text=text)], meta={"rows": total})


def _text(path: Path, ext: str) -> DocResult:
    text = path.read_bytes().decode("utf-8", errors="replace")
    if ext in CODE_EXT or ext in (".json", ".yaml", ".yml", ".toml", ".xml", ".sql"):
        lang = ext.lstrip(".")
        text = f"```{lang}\n{text}\n```"
        fmt = "code"
    else:
        fmt = "markdown" if ext in (".md", ".markdown") else "text"
    return DocResult(format=fmt, sections=[DocSection(label="", text=text)])


def _md_table(rows: list[list[str]]) -> str:
    if not rows:
        return ""
    width = max(len(r) for r in rows)
    norm = [[(c or "").replace("|", "\\|").replace("\n", " ") for c in r] + [""] * (width - len(r))
            for r in rows]
    head, *body = norm
    lines = ["| " + " | ".join(head) + " |", "|" + "---|" * width]
    lines += ["| " + " | ".join(r) + " |" for r in body]
    return "\n".join(lines)
