"""Image normalization, detection, frame planning, documents and cache."""

from __future__ import annotations

import csv
from pathlib import Path

import pytest
from PIL import Image

from media.cache import MediaCache
from media.detect import detect_kind
from media.documents import extract_document
from media.images import normalize_image
from media.types import Attachment, ContextBlock, Kind
from media.video import plan_frames

FIX = Path(__file__).resolve().parent.parent / "fixtures"


# --- images -----------------------------------------------------------------------------
def _jpeg_with_exif(path: Path, orientation: int = 6) -> None:
    im = Image.new("RGB", (400, 200), "red")
    exif = Image.Exif()
    exif[0x0112] = orientation          # Orientation: rotate 90 CW
    exif[0x010F] = "FakeCam"            # Make
    exif[0x8825] = {1: "N", 2: (37.0, 46.0, 30.0)}  # GPS IFD
    im.save(path, "JPEG", exif=exif.tobytes())


def test_exif_orientation_applied_then_stripped(tmp_path):
    src = tmp_path / "photo.jpg"
    _jpeg_with_exif(src)
    info = normalize_image(src, tmp_path / "out", max_side=1024)
    with Image.open(info.path) as out:
        assert out.size == (200, 400)  # rotated by orientation 6
        assert not out.getexif()        # no EXIF (incl. GPS) left
        assert "exif" not in out.info
    assert info.format == "jpeg"


def test_downscale_keeps_aspect(tmp_path):
    src = tmp_path / "big.png"
    Image.new("RGB", (3000, 1500), "blue").save(src)
    info = normalize_image(src, tmp_path / "out", max_side=896)
    assert (info.width, info.height) == (896, 448)
    assert info.format == "png"


def test_webp_and_alpha(tmp_path):
    src = tmp_path / "a.webp"
    Image.new("RGBA", (100, 100), (0, 0, 0, 0)).save(src, "WEBP")
    info = normalize_image(src, tmp_path / "out", max_side=896)
    assert info.format == "png"  # alpha preserved in PNG


def test_heic(tmp_path):
    pillow_heif = pytest.importorskip("pillow_heif")
    src = tmp_path / "IMG_0001.HEIC"
    heif = pillow_heif.from_pillow(Image.new("RGB", (640, 480), "green"))
    heif.save(src, quality=80)
    att = Attachment(src)
    assert detect_kind(att) is Kind.IMAGE
    info = normalize_image(src, tmp_path / "out", max_side=320)
    assert info.source_format in ("HEIF", "HEIC") and info.width == 320


# --- detection --------------------------------------------------------------------------
@pytest.mark.parametrize("name,source,kind", [
    ("chart.png", "upload", Kind.IMAGE),
    ("chart.png", "paste", Kind.SCREENSHOT),
    ("error_dialog.png", "screenshot", Kind.SCREENSHOT),
    ("dictation.wav", "upload", Kind.AUDIO),
    ("voice_question.m4a", "upload", Kind.AUDIO),
    ("tour.mp4", "upload", Kind.VIDEO),
])
def test_detect_fixtures(name, source, kind):
    assert detect_kind(Attachment(FIX / name, source=source)) is kind


def test_detect_screenshot_by_name(tmp_path):
    src = tmp_path / "Screenshot 2026-09-24 at 10.00.00.png"
    Image.new("RGB", (10, 10)).save(src)
    assert detect_kind(Attachment(src)) is Kind.SCREENSHOT


def test_detect_audio_webm_by_mime(tmp_path):
    src = tmp_path / "voice.webm"
    src.write_bytes(b"\x1a\x45\xdf\xa3" + b"\x00" * 64)
    assert detect_kind(Attachment(src, mime="audio/webm;codecs=opus")) is Kind.AUDIO


def test_detect_code_and_binary(tmp_path):
    code = tmp_path / "main.rs"
    code.write_text("fn main() {}")
    blob = tmp_path / "blob.bin"
    blob.write_bytes(b"\x00\x01\x02" * 100)
    assert detect_kind(Attachment(code)) is Kind.DOCUMENT
    assert detect_kind(Attachment(blob)) is Kind.UNKNOWN


# --- frame planning ---------------------------------------------------------------------
def test_plan_respects_budget_and_order():
    plan = plan_frames(600, [10, 50, 51, 300], budget=12)
    times = [t for t, _ in plan]
    assert len(plan) <= 12 and times == sorted(times)
    assert times[0] <= 0.5
    assert all(0 <= t < 600 for t in times)


def test_plan_includes_scene_changes_after_settle():
    plan = plan_frames(75, [20.0, 40.0, 60.0], budget=16)
    scene_times = [t for t, r in plan if r == "scene"]
    assert scene_times == [20.4, 40.4, 60.4]


def test_plan_caps_scene_share_and_keeps_coverage():
    scenes = [i * 0.5 for i in range(1, 400)]  # flashy video: a cut every half second
    plan = plan_frames(200, scenes, budget=10)
    assert len(plan) <= 10
    reasons = [r for _, r in plan]
    assert reasons.count("scene") <= 6 and "uniform" in reasons


def test_plan_short_clip_uses_few_frames():
    assert len(plan_frames(6, [], budget=32)) == 4
    assert plan_frames(0, [], budget=8) == []


# --- documents --------------------------------------------------------------------------
def test_docx(tmp_path):
    import docx

    d = docx.Document()
    d.add_heading("Quarterly plan", level=1)
    d.add_paragraph("Ship the widget by Friday.")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Owner", "Task"
    t.cell(1, 0).text, t.cell(1, 1).text = "Ana", "QA"
    f = tmp_path / "plan.docx"
    d.save(f)
    text = extract_document(f, f.name).text
    assert "# Quarterly plan" in text and "Ship the widget" in text and "| Ana | QA |" in text


def test_pptx(tmp_path):
    from pptx import Presentation

    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Roadmap"
    s.placeholders[1].text = "Launch in Q3"
    s.notes_slide.notes_text_frame.text = "Mention the budget"
    f = tmp_path / "deck.pptx"
    prs.save(f)
    res = extract_document(f, f.name)
    assert res.meta["slides"] == 1
    assert "slide 1" in res.text and "Launch in Q3" in res.text and "Mention the budget" in res.text


def test_xlsx_and_csv(tmp_path):
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Revenue"
    ws.append(["Quarter", "USD"])
    ws.append(["Q1", 42])
    f = tmp_path / "r.xlsx"
    wb.save(f)
    assert "| Q1 | 42 |" in extract_document(f, f.name).text

    c = tmp_path / "r.csv"
    with open(c, "w", newline="") as fh:
        csv.writer(fh).writerows([["a", "b"], ["1", "2"]])
    assert "| 1 | 2 |" in extract_document(c, c.name).text


def test_pdf_text_and_scanned_page(tmp_path):
    import fitz

    doc = fitz.open()
    p1 = doc.new_page()
    p1.insert_text((72, 72), "Invoice total: 1,234.56 EUR")
    p2 = doc.new_page()  # "scanned" page: an image, no text layer
    img = tmp_path / "scan.png"
    Image.new("RGB", (400, 200), "white").save(img)
    p2.insert_image(p2.rect, filename=str(img))
    f = tmp_path / "invoice.pdf"
    doc.save(f)
    res = extract_document(f, f.name, render_dir=tmp_path / "pages")
    assert "1,234.56" in res.sections[0].text and res.sections[0].image_path is None
    assert res.sections[1].image_path and Path(res.sections[1].image_path).exists()
    assert res.meta["scanned_pages"] == 1


def test_code_is_fenced(tmp_path):
    f = tmp_path / "app.py"
    f.write_text("print('hi')")
    assert extract_document(f, f.name).text.startswith("```py")


# --- cache / envelope -------------------------------------------------------------------
def test_cache_roundtrip_and_params(tmp_path):
    c = MediaCache(tmp_path)
    work = c.scratch("ab" * 32)
    (work / "x.txt").write_text("payload")
    entry = c.put("ab" * 32, "stage", {"k": 1}, {"ok": True}, files_from=work)
    assert c.get("ab" * 32, "stage", {"k": 1}) == {"ok": True}
    assert (entry / "x.txt").read_text() == "payload"
    assert c.get("ab" * 32, "stage", {"k": 2}) is None


def test_context_block_escapes_attrs():
    b = ContextBlock('a"b.txt', "document", "text")
    assert 'name="a&quot;b.txt"' in b.render()
